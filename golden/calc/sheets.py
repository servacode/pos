"""The results sheet (23 › «ورقة النتائج») for the owner and the accountant: results.xlsx and RESULTS.md."""

from pathlib import Path

from catalogue import BUSINESS_DATE, PRODUCT_BY_SKU, RATE_AFTERNOON, RATE_MORNING, USERS
from money import fmt, fmt_qty

METHOD_AR = {"cash": "نقداً", "card": "بطاقة", "wallet": "محفظة", "transfer": "حوالة", "credit": "آجل"}
CUR_AR = {"SYP": "ل.س", "USD": "$"}
KIND_AR = {"pay_in": "إيداع في الدرج", "pay_out": "سحب من الدرج", "expense": "مصروف من الدرج", "safe_drop": "تحويل للخزنة"}
DEV_AR = {"D1": "الجهاز 001", "D2": "الجهاز 002"}


def user_name(ref: str | None) -> str:
    if not ref:
        return ""
    return USERS[ref.split(":")[1]]["fullName"]


def money(minor: int, currency: str = "SYP") -> str:
    return f"{fmt(minor)} {CUR_AR[currency]}"


def discount_text(d: dict | None) -> str:
    if not d:
        return ""
    base = f"{d['bp'] / 100:g}%" if d["type"] == "percent" else fmt(d["minor"])
    if d.get("source") == "price_override":
        base = f"تعديل سعر ({base})"
    if d.get("approval"):
        base += " بموافقة المدير"
    return base


def write_all(day: Path, scenario: dict, expected: dict) -> None:
    write_xlsx(day / "results.xlsx", scenario, expected)
    (day / "RESULTS.md").write_text(results_md(scenario, expected), encoding="utf-8")


# ─── xlsx ─────────────────────────────────────────────────────────────────────


def write_xlsx(path: Path, scenario: dict, expected: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title: str, header: list[str], rows: list[list]):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
        return ws

    wb.remove(wb.active)
    d = expected["dayReport"]
    sheet(
        "الملخص",
        ["البند", "القيمة (ل.س)"],
        [
            ["تاريخ اليوم", BUSINESS_DATE],
            ["عدد الفواتير", d["invoiceCount"]],
            ["الإجمالي بأسعار الرف", d["grossMinor"] / 100],
            ["الخصومات بأسعار الرف", d["discountMinor"] / 100],
            ["المبيعات قبل الضريبة (4100)", d["grossExTaxMinor"] / 100],
            ["الحسميات قبل الضريبة (4200)", d["discountExTaxMinor"] / 100],
            ["الصافي قبل الضريبة", d["netExTaxMinor"] / 100],
            ["رسم الإنفاق الاستهلاكي (2210)", d["taxMinor"] / 100],
            ["مجموع الفواتير", d["totalMinor"] / 100],
            ["الإكراميات (2550)", d["tipsMinor"] / 100],
            ["فروقات التقريب (6400، موجب = ربح)", d["cashRoundingMinor"] / 100],
            ["سعر الصرف صباحاً", RATE_MORNING["rate"]],
            ["سعر الصرف من 14:00", RATE_AFTERNOON["rate"]],
        ],
    )
    sheet(
        "الفواتير",
        ["الرقم", "الجهاز", "الوقت", "دون إنترنت", "السعر", "الإجمالي", "الخصم", "الضريبة", "المجموع", "الإكرامية", "التقريب", "الباقي", "عملة الباقي"],
        [
            [
                i["number"], DEV_AR[i["device"]], i["deviceTime"], "نعم" if i["offline"] else "", i["rate"],
                i["totals"]["grossMinor"] / 100, i["totals"]["discountMinor"] / 100, i["totals"]["taxMinor"] / 100,
                i["totals"]["totalMinor"] / 100, i["tipMinor"] / 100,
                (i["cashRounding"]["roundingBaseMinor"] / 100) if i["cashRounding"] else 0,
                (i["change"]["amountMinor"] / 100) if i["change"] else 0,
                i["change"]["currency"] if i["change"] else "",
            ]
            for i in expected["invoices"]
        ],
    )
    sheet(
        "البنود",
        ["الفاتورة", "السطر", "الصنف", "الكمية", "سعر الوحدة", "الإجمالي", "خصم البند", "حصة خصم الفاتورة", "الصافي", "الضريبة", "الإجمالي قبل الضريبة", "الخصم قبل الضريبة", "المجموع", "خروج من المخزون"],
        [
            [
                i["number"], ln["lineNo"], ln["name"], fmt_qty(ln["qty"]), ln["unitPriceMinor"] / 100, ln["grossMinor"] / 100,
                ln["lineDiscountMinor"] / 100, ln["invoiceDiscountShareMinor"] / 100, ln["netMinor"] / 100, ln["taxMinor"] / 100,
                ln["grossExTaxMinor"] / 100, ln["discountExTaxMinor"] / 100, ln["totalMinor"] / 100, fmt_qty(ln["baseQty"]),
            ]
            for i in expected["invoices"]
            for ln in i["lines"]
        ],
    )
    sheet(
        "الدفعات",
        ["الفاتورة", "رقم الدفعة", "الطريقة", "العملة", "المبلغ", "السعر", "بالليرة"],
        [
            [i["number"], p["paymentNo"], METHOD_AR[p["method"]], p["currency"], p["amountMinor"] / 100, p["rate"]["rate"] if p["rate"] else "", p["baseAmountMinor"] / 100]
            for i in expected["invoices"]
            for p in i["payments"]
        ],
    )
    sheet(
        "الورديات",
        ["الوردية", "العملة", "الافتتاحي", "المقبوض نقداً", "الباقي المدفوع", "إيداعات", "سحوبات ومصروفات", "للخزنة", "المتوقع", "المعدود", "الفرق", "الفرق بالليرة"],
        [
            [
                s["number"], c["currency"], c["openingMinor"] / 100, c["cashSalesMinor"] / 100, c["changeGivenMinor"] / 100, c["payInMinor"] / 100,
                c["payOutMinor"] / 100, c["dropsMinor"] / 100, c["expectedMinor"] / 100, c["countedMinor"] / 100, c["differenceMinor"] / 100,
                c["differenceBaseMinor"] / 100,
            ]
            for s in expected["shifts"]
            for c in s["byCurrency"]
        ],
    )
    sheet(
        "المخزون",
        ["الرمز", "الصنف", "أول اليوم", "المباع", "آخر اليوم"],
        [[r["sku"], PRODUCT_BY_SKU[r["sku"]]["nameAr"], fmt_qty(r["openingQty"]), fmt_qty(r["soldQty"]), fmt_qty(r["closingQty"])] for r in expected["stock"]],
    )
    sheet(
        "القيود",
        ["المصدر", "الرقم", "الحساب", "اسم الحساب", "مدين", "دائن", "العملة", "المبلغ بعملته", "السعر"],
        [
            [e["source"]["type"], e["source"]["number"], ln["accountCode"], ln["accountName"], ln["debitMinor"] / 100, ln["creditMinor"] / 100, ln["currency"], ln["amountMinor"] / 100, ln["rate"] or ""]
            for e in expected["journal"]
            for ln in e["lines"]
        ],
    )
    tb = expected["trialBalance"]
    ws = sheet(
        "ميزان المراجعة",
        ["الحساب", "الاسم", "مدين", "دائن", "الرصيد", "الطبيعة"],
        [[r["accountCode"], r["accountName"], r["debitMinor"] / 100, r["creditMinor"] / 100, r["closingMinor"] / 100, "مدين" if r["normalSide"] == "debit" else "دائن"] for r in tb["rows"]]
        + [["", "المجموع", tb["totals"]["debitMinor"] / 100, tb["totals"]["creditMinor"] / 100, "", "متوازن" if tb["balanced"] else "غير متوازن"]],
    )
    ws.cell(row=ws.max_row, column=2).font = bold
    sheet(
        "الاعتماد",
        ["الدور", "الاسم", "التوقيع", "التاريخ"],
        [["المالك", "", "", ""], ["المحاسب", "", "", ""]],
    )
    # fixed workbook metadata so the file does not change on every run
    from datetime import datetime

    fixed = datetime(2026, 10, 30, 0, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "golden/calc"
    # openpyxl stamps every zip entry with the current time: rewrite them with a fixed one
    import io
    import zipfile

    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed_info = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 30, 0, 0, 0))
            fixed_info.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl writes «modified» as now
                text = data.decode("utf-8")
                start = text.index("<dcterms:modified")
                start = text.index(">", start) + 1
                end = text.index("</dcterms:modified>", start)
                data = (text[:start] + "2026-10-30T00:00:00Z" + text[end:]).encode("utf-8")
            dst.writestr(fixed_info, data)
    path.write_bytes(out.getvalue())


# ─── RESULTS.md ───────────────────────────────────────────────────────────────


def results_md(scenario: dict, expected: dict) -> str:
    d = expected["dayReport"]
    out = []
    w = out.append
    w("# ورقة نتائج اليوم الذهبي — المرحلة 1")
    w("")
    w(f"«ماركت الشام»، فرع 001، يوم {BUSINESS_DATE}، جهازا كاشير (001 و002) ووردية على كل جهاز.")
    w(f"سعر الدولار {RATE_MORNING['rate']} ل.س حتى الساعة 14:00، ثم {RATE_AFTERNOON['rate']} ل.س. الأسعار تشمل الرسم.")
    w("")
    w("هذه الورقة محسوبة بحاسبة مستقلة عن البرنامج (`golden/calc`). **لا تُعتمد قبل توقيع المالك ومحاسب** (تبويب الاختبار والقبول).")
    w("")
    w("## الملخص")
    w("")
    w("| البند | القيمة |")
    w("| --- | --- |")
    rows = [
        ("عدد الفواتير", f"{d['invoiceCount']}"),
        ("منها دون إنترنت ومُرسلة مرتين", f"{sum(1 for i in expected['invoices'] if i['offline'])}"),
        ("الإجمالي بأسعار الرف", money(d["grossMinor"])),
        ("الخصومات", money(d["discountMinor"])),
        ("مجموع الفواتير", money(d["totalMinor"])),
        ("منها رسم الإنفاق الاستهلاكي", money(d["taxMinor"])),
        ("المبيعات قبل الرسم والخصم (4100)", money(d["grossExTaxMinor"])),
        ("الحسميات قبل الرسم (4200)", money(d["discountExTaxMinor"])),
        ("الإكراميات (ليست مبيعات)", money(d["tipsMinor"])),
        ("فروقات تقريب النقد (موجب = لصالح المحل)", money(d["cashRoundingMinor"])),
    ]
    for k, v in rows:
        w(f"| {k} | {v} |")
    w("")
    w("### حسب طريقة الدفع")
    w("")
    w("| الطريقة | العملة | عدد الفواتير | المبلغ بعملته (بعد الباقي) | بالليرة |")
    w("| --- | --- | --- | --- | --- |")
    for r in d["byMethod"]:
        w(f"| {METHOD_AR[r['method']]} | {CUR_AR[r['currency']]} | {r['count']} | {fmt(r['amountMinor'])} | {fmt(r['baseAmountMinor'])} |")
    w("")
    w("## أول عشر فواتير بالتفصيل")
    w("")
    for inv in expected["invoices"][:10]:
        w(f"### {inv['number']} — {DEV_AR[inv['device']]}، السعر {inv['rate']}" + (" — دون إنترنت" if inv["offline"] else ""))
        w("")
        w("| # | الصنف | الكمية × السعر | الإجمالي | خصم البند | حصة خصم الفاتورة | الصافي | منه الرسم |")
        w("| --- | --- | --- | --- | --- | --- | --- | --- |")
        for ln in inv["lines"]:
            ld = fmt(ln["lineDiscountMinor"])
            if ln["discount"]:
                ld += f" ({discount_text(ln['discount'])})"
            w(
                f"| {ln['lineNo']} | {ln['name']} | {fmt_qty(ln['qty'])} × {fmt(ln['unitPriceMinor'])} | {fmt(ln['grossMinor'])} | {ld} | "
                f"{fmt(ln['invoiceDiscountShareMinor'])} | {fmt(ln['netMinor'])} | {fmt(ln['taxMinor'])} |"
            )
        t = inv["totals"]
        w(f"| | **المجموع** | | **{fmt(t['grossMinor'])}** | **{fmt(t['lineDiscountMinor'])}** | **{fmt(t['invoiceDiscountMinor'])}** | **{fmt(t['totalMinor'])}** | **{fmt(t['taxMinor'])}** |")
        w("")
        if inv["invoiceDiscount"]:
            w(f"- خصم الفاتورة: {discount_text(inv['invoiceDiscount'])} = {fmt(t['invoiceDiscountMinor'])}، يُوزَّع على البنود بنسبة صافي كل بند.")
        if inv["tipMinor"]:
            w(f"- إكرامية: {money(inv['tipMinor'])}، فالمستحق {money(t['totalMinor'] + inv['tipMinor'])}.")
        for p in inv["payments"]:
            rate = f" بسعر {p['rate']['rate']} = {money(p['baseAmountMinor'])}" if p["rate"] else ""
            w(f"- دفع {METHOD_AR[p['method']]}: {money(p['amountMinor'], p['currency'])}{rate}.")
        cr = inv["cashRounding"]
        if cr:
            w(
                f"- النقد المستحق بـ{CUR_AR[cr['currency']]} قبل التقريب {fmt(cr['exactMinor'])}، وبعد التقريب لأقرب {fmt(cr['stepMinor'])}: "
                f"{fmt(cr['dueMinor'])}؛ فرق التقريب {money(cr['roundingBaseMinor'])}."
            )
        if inv["change"]:
            w(f"- الباقي للزبون: {money(inv['change']['amountMinor'], inv['change']['currency'])}.")
        w("- القيد: " + "، ".join(
            f"{'مدين' if ln['debitMinor'] else 'دائن'} {ln['accountCode']} {ln['accountName']}"
            + (f" ({CUR_AR[ln['currency']]})" if ln["accountCode"] == "1110" else "")
            + f" {fmt(ln['debitMinor'] or ln['creditMinor'])}"
            for ln in inv["journal"]
        ) + ".")
        w("")
    w("## النقد في كل وردية")
    w("")
    for s in expected["shifts"]:
        w(f"### الوردية {s['number']} — {DEV_AR[s['device']]}، الكاشير {user_name(s['openedBy'])}")
        w("")
        w(f"{s['salesCount']} فاتورة بمجموع {money(s['salesTotalMinor'])}.")
        w("")
        w("| العملة | الافتتاحي | المقبوض نقداً | الباقي المدفوع | إيداعات | سحوبات ومصروفات | للخزنة | المتوقع | المعدود | الفرق |")
        w("| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |")
        for c in s["byCurrency"]:
            diff = c["differenceMinor"]
            label = "زيادة" if diff > 0 else ("عجز" if diff < 0 else "لا فرق")
            w(
                f"| {CUR_AR[c['currency']]} | {fmt(c['openingMinor'])} | {fmt(c['cashSalesMinor'])} | {fmt(c['changeGivenMinor'])} | {fmt(c['payInMinor'])} | "
                f"{fmt(c['payOutMinor'])} | {fmt(c['dropsMinor'])} | {fmt(c['expectedMinor'])} | {fmt(c['countedMinor'])} | {label} {fmt(abs(diff))} |"
            )
        w("")
        for m in s["movements"]:
            w(f"- {m['number']}: {KIND_AR[m['kind']]} {money(m['amountMinor'], m['currency'])} — «{m['reason']}»، بموافقة {user_name(m['approvedBy'])}.")
        w("")
    w("## ميزان المراجعة لليوم")
    w("")
    w("الأرقام بالليرة. الرصيد حسب طبيعة الحساب: «له» للدائن و«عليه» للمدين. لا رصيد افتتاحياً: المتجر جديد، والرصيد الافتتاحي للدرج خارج القيود (انظر INTERPRETATIONS.md).")
    w("")
    w("| الحساب | الاسم | مدين | دائن | الرصيد |")
    w("| --- | --- | --- | --- | --- |")
    for r in expected["trialBalance"]["rows"]:
        side = "عليه" if (r["normalSide"] == "debit") == (r["closingMinor"] >= 0) else "له"
        w(f"| {r['accountCode']} | {r['accountName']} | {fmt(r['debitMinor'])} | {fmt(r['creditMinor'])} | {fmt(abs(r['closingMinor']))} {side} |")
    t = expected["trialBalance"]["totals"]
    w(f"| | **المجموع** | **{fmt(t['debitMinor'])}** | **{fmt(t['creditMinor'])}** | {'متوازن' if expected['trialBalance']['balanced'] else 'غير متوازن'} |")
    w("")
    w("## المخزون آخر اليوم")
    w("")
    w("| الرمز | الصنف | أول اليوم | المباع | آخر اليوم |")
    w("| --- | --- | --- | --- | --- |")
    for r in expected["stock"]:
        w(f"| {r['sku']} | {PRODUCT_BY_SKU[r['sku']]['nameAr']} | {fmt_qty(r['openingQty'])} | {fmt_qty(r['soldQty'])} | {fmt_qty(r['closingQty'])} |")
    w("")
    w("## شاشة الدليل المتوقعة")
    w("")
    w("| الجهاز | أُرسل | قُبل | مكرر | رُفض | محجوز | آخر رقم أعلنه = آخر رقم وصل |")
    w("| --- | --- | --- | --- | --- | --- | --- |")
    for dev, e in expected["evidence"].items():
        w(f"| {e['deviceCode']} | {e['sent']} | {e['accepted']} | {e['duplicate']} | {e['rejected']} | {e['quarantined']} | {', '.join(e['lastAnnouncedNumber'].values())} |")
    w("")
    w("## الاعتماد")
    w("")
    w("- [ ] المالك: ____________________ التاريخ: __________")
    w("- [ ] المحاسب: ____________________ التاريخ: __________")
    w("")
    return "\n".join(out)
