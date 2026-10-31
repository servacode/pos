"""Checks the independent calculator against the worked examples of spec 10 and 23, literally.

    python3 -m unittest discover -s golden/calc -p 'test_*.py'
"""

import json
import random
import unittest
from pathlib import Path

from ledger import balanced, post_sale
from money import convert, round_to_step
from pricing import CartLine, allocate, price_invoice
from settle import settle

DAY = Path(__file__).resolve().parent.parent / "day1"
R130 = {"unit": "USD", "quote": "SYP", "rate": "130"}
STEPS = {"SYP": 1000, "USD": 25}


class SpecTen(unittest.TestCase):
    def test_worked_invoice_exclusive_15(self):
        """10 › حساب الفاتورة, the table: milk, bread, rice, 0.40 off the rice, 1.00 off the invoice."""
        inv = price_invoice(
            [
                CartLine(1, 125, 2000, 1500),
                CartLine(2, 75, 3000, 1500),
                CartLine(3, 890, 1000, 1500, {"type": "amount", "minor": 40}),
            ],
            {"type": "amount", "minor": 100},
            tax_inclusive=False,
        )
        self.assertEqual([l.share for l in inv.lines], [19, 17, 64])
        self.assertEqual([l.net for l in inv.lines], [231, 208, 786])
        self.assertEqual([l.tax for l in inv.lines], [35, 31, 118])
        self.assertEqual([l.total for l in inv.lines], [266, 239, 904])
        self.assertEqual(inv.totals["grossMinor"], 1365)
        self.assertEqual(inv.totals["discountMinor"], 140)
        self.assertEqual(inv.totals["netMinor"], 1225)
        self.assertEqual(inv.totals["taxMinor"], 184)
        self.assertEqual(inv.totals["totalMinor"], 1409)
        # «دفع الزبون 20.00 نقداً، فالباقي 5.91» (no cash step in this example)
        st = settle("SYP", 1409, 0, [{"method": "cash", "currency": "SYP", "amountMinor": 2000}], "SYP", {}, {})
        self.assertEqual(st.change["amountMinor"], 591)
        # and its entry: 1110 14.09 + 4200 1.40 = 4100 13.65 + tax 1.84
        sale = {
            "payments": st.payments,
            "change": st.change,
            "totals": inv.totals,
            "lines": [{"taxRateId": "t", "priced": l.as_payload()} for l in inv.lines],
            "tipMinor": 0,
            "cashRounding": st.cash_rounding,
        }
        lines = post_sale(sale, {"t": "sales"})
        got = {(l["accountCode"], l["debitMinor"], l["creditMinor"]) for l in lines}
        self.assertEqual(got, {("1110", 1409, 0), ("4200", 140, 0), ("4100", 0, 1365), ("2200", 0, 184)})

    def test_inclusive_price_example(self):
        """10: 1,150 including 15%, a 150 discount: 1,000 / 870 / 130 / 130, both sides 1,130 (minor units)."""
        inv = price_invoice([CartLine(1, 1_150, 1000, 1500, {"type": "amount", "minor": 150})], None, True)
        line = inv.lines[0]
        self.assertEqual(line.gross_ex_tax, 1_000)
        self.assertEqual(line.net_ex_tax, 870)
        self.assertEqual(line.tax, 130)
        self.assertEqual(line.discount_ex_tax, 130)
        self.assertEqual(line.total, 1_000)

    def test_cash_rounding_example(self):
        """10: 12,347.50 with a 10-pound step: cash pays 12,350 and 2.50 is credited to 6400; credit stays exact."""
        st = settle("SYP", 1_234_750, 0, [{"method": "cash", "currency": "SYP", "amountMinor": 1_235_000}], "SYP", {}, STEPS)
        self.assertEqual(st.cash_rounding["roundingBaseMinor"], 250)
        self.assertIsNone(st.change)
        self.assertEqual(round_to_step(1_234_750, 1000), 1_235_000)

    def test_rate_divides(self):
        """10: 162,500 SYP at 13,000 = 12.50 $ exactly, not 12.51."""
        self.assertEqual(convert(16_250_000, "SYP", "USD", {"unit": "USD", "quote": "SYP", "rate": "13000"}), 1250)

    def test_half_away_from_zero(self):
        self.assertEqual(round_to_step(2500, 1000), 3000)
        self.assertEqual(round_to_step(-2500, 1000), -3000)
        self.assertEqual(round_to_step(25 * 1 + 12, 25), 25)

    def test_tie_goes_to_smaller_line(self):
        """23 › الحالات الحدّية: «تعادل في توزيع الخصم», ties by line_no (10 › step 3)."""
        self.assertEqual(allocate(5, [(1, 60_000), (2, 60_000)]), {1: 3, 2: 2})
        self.assertEqual(allocate(5, [(2, 60_000), (1, 60_000)]), {1: 3, 2: 2})


class SpecTwentyThree(unittest.TestCase):
    def test_first_sale(self):
        """23 › 3: 1,290.00; 9 $ at 130 is 1,170.00; 120.00 left; a 200 note; 80.00 back."""
        inv = price_invoice(
            [
                CartLine(1, 5_000, 3000, 0),
                CartLine(2, 32_500, 1000, 0),
                CartLine(3, 60_000, 1000, 0, {"type": "percent", "bp": 1000}),
                CartLine(4, 27_500, 1000, 0),
            ],
            None,
            True,
        )
        self.assertEqual(inv.lines[2].total, 54_000)
        self.assertEqual(inv.totals["totalMinor"], 129_000)
        st = settle(
            "SYP",
            129_000,
            0,
            [{"method": "cash", "currency": "USD", "amountMinor": 900}, {"method": "cash", "currency": "SYP", "amountMinor": 20_000}],
            "SYP",
            {"USD": R130},
            STEPS,
        )
        self.assertEqual(st.payments[0]["baseAmountMinor"], 117_000)
        self.assertEqual(st.cash_rounding["exactMinor"], 12_000)
        self.assertEqual(st.change["amountMinor"], 8_000)
        # «وبالدولار 10.38» for 1,350.00
        self.assertEqual(convert(135_000, "SYP", "USD", R130), 1038)


class Properties(unittest.TestCase):
    """Random sales: the payload balance of documents.ts holds and every entry balances."""

    def test_random_sales_balance(self):
        rng = random.Random(7)
        for _ in range(3000):
            lines = [
                CartLine(
                    i,
                    rng.randint(1, 500_000),
                    rng.choice([250, 333, 1000, 2000, 1250]),
                    rng.choice([0, 0, 1500]),
                    rng.choice([None, {"type": "percent", "bp": rng.randint(0, 10_000)}, {"type": "amount", "minor": rng.randint(0, 10_000)}]),
                )
                for i in range(1, rng.randint(1, 8))
            ]
            inv = price_invoice(lines, rng.choice([None, {"type": "amount", "minor": rng.randint(0, 50_000)}]), rng.random() < 0.5)
            total = inv.totals["totalMinor"]
            self.assertEqual(sum(l.share for l in inv.lines), inv.invoice_discount)
            self.assertEqual(inv.totals["grossExTaxMinor"] - inv.totals["discountExTaxMinor"] + inv.totals["taxMinor"], total)
            if total == 0:
                continue
            tip = rng.choice([0, 0, 5000])
            rate = {"unit": "USD", "quote": "SYP", "rate": rng.choice(["130", "132", "131.5"])}
            usd = rng.randint(0, 3) * 100
            s = rng.choice(["SYP", "USD"])
            due = total + tip
            remaining = due - convert(usd, "USD", "SYP", rate)
            if s == "SYP":
                tender = [{"method": "cash", "currency": "USD", "amountMinor": usd}] if usd else []
                tender.append({"method": "cash", "currency": "SYP", "amountMinor": max(1, round_to_step(remaining, 1000) + 50_000)})
            else:
                need = round_to_step(convert(due, "SYP", "USD", rate), 25)
                tender = [{"method": "cash", "currency": "USD", "amountMinor": max(25, need + 100)}]
            st = settle("SYP", total, tip, tender, s, {"USD": rate}, STEPS)
            paid = sum(p["baseAmountMinor"] for p in st.payments) - (st.change["baseAmountMinor"] if st.change else 0)
            self.assertEqual(paid, total + tip + st.cash_rounding["roundingBaseMinor"])
            cr = st.cash_rounding
            self.assertEqual(cr["dueMinor"] % STEPS[cr["currency"]], 0)
            self.assertLessEqual(abs(cr["roundingMinor"]) * 2, STEPS[cr["currency"]])
            sale = {
                "payments": st.payments,
                "change": st.change,
                "totals": inv.totals,
                "lines": [{"taxRateId": "t", "priced": l.as_payload()} for l in inv.lines],
                "tipMinor": tip,
                "cashRounding": st.cash_rounding,
            }
            self.assertTrue(balanced(post_sale(sale, {"t": "consumption"})))

    def test_line_order_never_changes_totals(self):
        rng = random.Random(11)
        for _ in range(500):
            lines = [CartLine(i, rng.randint(1, 90_000), rng.choice([333, 1000, 3000]), rng.choice([0, 1500])) for i in range(1, 7)]
            disc = {"type": "amount", "minor": rng.randint(0, 20_000)}
            a = price_invoice(lines, disc, True)
            shuffled = lines[:]
            rng.shuffle(shuffled)
            b = price_invoice(shuffled, disc, True)
            self.assertEqual(a.totals, b.totals)
            self.assertEqual([l.as_payload() for l in a.lines], [l.as_payload() for l in b.lines])


class GoldenFiles(unittest.TestCase):
    def test_expected_is_balanced_and_complete(self):
        e = json.loads((DAY / "expected.json").read_text(encoding="utf-8"))
        s = json.loads((DAY / "scenario.json").read_text(encoding="utf-8"))
        self.assertEqual(e["dayReport"]["invoiceCount"], 200)
        self.assertEqual(sum(1 for i in e["invoices"] if i["offline"]), 50)
        self.assertTrue(e["trialBalance"]["balanced"])
        for entry in e["journal"]:
            self.assertTrue(balanced(entry["lines"]))
        first = e["invoices"][0]
        self.assertEqual(first["number"], "INV-001-001-0000001")
        self.assertEqual(first["totals"]["totalMinor"], 129_000)
        self.assertEqual(first["change"], {"currency": "SYP", "amountMinor": 8_000, "rate": None, "baseAmountMinor": 8_000})
        diffs = sorted(c["differenceMinor"] for sh in e["shifts"] for c in sh["byCurrency"] if c["currency"] == "SYP")
        self.assertEqual(diffs, [-50_000, 100_000])
        pushed = [i for b in s["pushPlan"] for i in b["documents"]]
        unique = {d["id"] for docs in s["documents"].values() for d in docs}
        self.assertEqual(set(pushed), unique)
        self.assertEqual(len(pushed), len(unique) + 50)
        self.assertTrue(all(len(b["documents"]) <= 50 for b in s["pushPlan"]))
        self.assertTrue(all(len(set(b["documents"])) == len(b["documents"]) for b in s["pushPlan"]))

    def test_sheets_follow_the_glossary(self):
        """Tab 14 «نقول ولا نقول», as in apps/web/src/test/glossary.ts: the sheet is read by a shopkeeper."""
        banned = ["عميل", "عملاء", "مستهلك", "PIN", "مزامنة", "متزامن", "Sync", "جلسة", "جلست", "Offline", "كلمة المرور", "Password"]
        from openpyxl import load_workbook

        texts = [(DAY / "RESULTS.md").read_text(encoding="utf-8")]
        wb = load_workbook(DAY / "results.xlsx")
        for ws in wb.worksheets:
            texts.append(ws.title)
            texts.extend(str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
        found = sorted({w for w in banned for t in texts if w in t})
        self.assertEqual(found, [])

    def test_regeneration_is_deterministic(self):
        import subprocess
        import sys

        before = {p.name: p.read_bytes() for p in DAY.iterdir() if p.suffix in (".json", ".md", ".xlsx")}
        subprocess.run([sys.executable, str(Path(__file__).with_name("build.py"))], check=True, capture_output=True)
        after = {p.name: p.read_bytes() for p in DAY.iterdir() if p.suffix in (".json", ".md", ".xlsx")}
        for name in ("scenario.json", "expected.json", "RESULTS.md"):
            self.assertEqual(before[name], after[name], name)


if __name__ == "__main__":
    unittest.main()
